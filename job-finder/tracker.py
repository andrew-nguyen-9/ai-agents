#!/usr/bin/env python3
"""
tracker.py — system of record for the Job Finding workflow.

Subcommands:
  init                      Create tracker.xlsx with headers (no-op if it exists).
  seed                      Seed rows from the existing jobs/ folders (via INDEX.md).
  add  <jobs.json>          Append found jobs (de-duped) + create jobs/<id>/ folders.
                            Pass "-" to read JSON from stdin.

The master file is <ROOT>/tracker.xlsx. ROOT is the "Job Finding" folder, inferred
as two levels up from this script (automation/job-finder/tracker.py).

A job record (for `add`) is a JSON object. Recognized fields:
  company, role, location, remote, pay_range, industry, ats, qualified,
  fit_rating, fit_reason, url, source_query, posting (full JD text)
Anything missing is left blank. `id` is auto-slugged from company+role if absent.
"""
import sys, os, json, re, datetime, contextlib
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

try:
    import fcntl  # POSIX advisory file locking (Linux/macOS)
except ImportError:
    fcntl = None


@contextlib.contextmanager
def file_lock(path):
    """Exclusive cross-process lock so parallel `add` calls can't clobber the xlsx.
    Falls back to a no-op if fcntl is unavailable."""
    if fcntl is None:
        yield
        return
    lock_path = str(path) + ".lock"
    f = open(lock_path, "w")
    try:
        fcntl.flock(f, fcntl.LOCK_EX)   # blocks until the lock is free
        yield
    finally:
        fcntl.flock(f, fcntl.LOCK_UN)
        f.close()

ROOT = Path(__file__).resolve().parents[2]          # the "Job Finding" folder
XLSX = ROOT / "tracker.xlsx"
JOBS = ROOT / "jobs"
SHEET = "jobs"

COLUMNS = [
    "id", "company", "role", "location", "remote", "pay_range", "industry",
    "ats", "qualified", "fit_rating", "fit_reason", "url", "source_query",
    "date_found", "decision", "status", "date_applied", "next_action", "notes",
]

def today():
    return datetime.date.today().isoformat()

def slugify(*parts):
    s = "-".join(p for p in parts if p)
    s = s.lower()
    s = re.sub(r"[^a-z0-9]+", "-", s)
    s = re.sub(r"-+", "-", s).strip("-")
    return s[:80]

# ---------- workbook helpers ----------
def load_or_create():
    if XLSX.exists():
        wb = openpyxl.load_workbook(XLSX)
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
        return wb, ws
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(COLUMNS)
    header_fill = PatternFill("solid", fgColor="1F2937")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    widths = {"id":28,"company":20,"role":34,"location":16,"remote":10,"pay_range":18,
              "industry":16,"ats":12,"qualified":10,"fit_rating":10,"fit_reason":40,
              "url":46,"source_query":34,"date_found":12,"decision":10,"status":16,
              "date_applied":12,"next_action":24,"notes":40}
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = widths.get(col, 16)
    return wb, ws

def existing_keys(ws):
    ids, urls = set(), set()
    idx = {c: i for i, c in enumerate(COLUMNS)}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[idx["id"]] in (None, ""):
            continue
        ids.add(str(row[idx["id"]]).strip())
        u = row[idx["url"]]
        if u:
            urls.add(str(u).strip())
    return ids, urls

def append_row(ws, rec):
    ws.append([rec.get(c, "") for c in COLUMNS])

# ---------- subcommands ----------
def cmd_init():
    wb, ws = load_or_create()
    wb.save(XLSX)
    print(f"tracker.xlsx ready ({ws.max_row-1} rows) at {XLSX}")

def parse_index_rows():
    """Pull (company, role, folder, status) from INDEX.md's application table."""
    index = ROOT / "INDEX.md"
    rows = []
    if not index.exists():
        return rows
    for line in index.read_text().splitlines():
        m = re.match(r"^\|\s*(.+?)\s*\|\s*(.+?)\s*\|\s*`jobs/([^`]+?)/?`\s*\|", line)
        if not m:
            continue
        company, role, folder = m.group(1), m.group(2), m.group(3)
        if company.lower() == "company":
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        status = cells[5] if len(cells) > 5 else ""
        status = "" if status in ("—", "-") else status
        rows.append((company, role, folder, status))
    return rows

def cmd_seed():
    wb, ws = load_or_create()
    ids, urls = existing_keys(ws)
    added = 0
    for company, role, folder, status in parse_index_rows():
        if folder in ids:
            continue
        has_cover = (JOBS / folder / "cover-letter.txt").exists()
        rec = {
            "id": folder, "company": company, "role": role,
            "status": status or ("draft" if has_cover else "Found"),
            "date_found": today(), "notes": "seeded from existing folder",
        }
        append_row(ws, rec)
        ids.add(folder)
        added += 1
    wb.save(XLSX)
    print(f"seed: added {added} existing applications ({ws.max_row-1} total rows)")

def is_qualified(job, min_fit=3):
    """Gate for saving a JD + folder. Qualified = explicit Yes/Stretch/Maybe,
    or (when unspecified) a fit_rating at/above min_fit. Explicit No/Skip = not."""
    q = str(job.get("qualified", "")).strip().lower()
    if q in ("no", "skip", "n"):
        return False
    if q in ("yes", "stretch", "maybe", "y"):
        return True
    try:
        return float(job.get("fit_rating", 0)) >= min_fit
    except (TypeError, ValueError):
        return False

def cmd_add(src):
    raw = sys.stdin.read() if src == "-" else Path(src).read_text()
    data = json.loads(raw)
    jobs = data if isinstance(data, list) else [data]
    added, low, skipped = [], [], []
    # Hold an exclusive lock across read->append->save so parallel agents (one per
    # job board) serialize their writes and each sees the others' just-added rows.
    with file_lock(XLSX):
        wb, ws = load_or_create()
        ids, urls = existing_keys(ws)
        for job in jobs:
            jid = job.get("id") or slugify(job.get("company",""), job.get("role",""))
            url = (job.get("url") or "").strip()
            if jid in ids or (url and url in urls):
                skipped.append(jid); continue
            qualified = is_qualified(job)
            rec = {c: job.get(c, "") for c in COLUMNS}
            rec["id"] = jid
            rec["date_found"] = today()
            rec["status"] = "Found" if qualified else "Skipped (low fit)"
            rec["decision"] = ""
            append_row(ws, rec)
            ids.add(jid)
            if url: urls.add(url)
            if qualified:
                # only qualified jobs get a folder + the saved JD
                folder = JOBS / jid
                folder.mkdir(parents=True, exist_ok=True)
                posting = job.get("posting", "")
                if posting:
                    (folder / "posting.txt").write_text(posting)
                meta = {k: job.get(k) for k in (
                    "company","role","location","remote","pay_range","industry","ats",
                    "url","source_query","fit_rating","fit_reason","qualified") if k in job}
                meta["id"] = jid
                meta["date_found"] = today()
                (folder / "meta.json").write_text(json.dumps(meta, indent=2))
                added.append(jid)
            else:
                low.append(jid)
        wb.save(XLSX)
    print(f"add: {len(added)} qualified (saved), {len(low)} low-fit (tracked only), "
          f"{len(skipped)} skipped (dupes)")
    if added:   print("  saved:   " + ", ".join(added))
    if low:     print("  low-fit: " + ", ".join(low))
    if skipped: print("  skipped: " + ", ".join(skipped))

def main():
    if len(sys.argv) < 2:
        print(__doc__); sys.exit(1)
    cmd = sys.argv[1]
    if cmd == "init": cmd_init()
    elif cmd == "seed": cmd_seed()
    elif cmd == "add":
        if len(sys.argv) < 3:
            print("usage: tracker.py add <jobs.json|->"); sys.exit(1)
        cmd_add(sys.argv[2])
    else:
        print(__doc__); sys.exit(1)

if __name__ == "__main__":
    main()
